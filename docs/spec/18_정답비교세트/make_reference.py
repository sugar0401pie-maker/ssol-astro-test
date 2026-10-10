"""
쏠 아스트로 하우스 — 출생차트 정답 비교 세트 생성기
독립 기준: skyfield + JPL DE421 (엔진은 astronomy-engine — 코드·천체력 모두 다름)

실행:  python make_reference.py
  1) 30건 테스트 출생정보 → zoneinfo로 UTC 변환
  2) skyfield/DE421로 10행성 경도·별자리·역행, ASC·MC, 홀사인 하우스 계산
  3) 정답세트.json 저장 → compare_engine.run() 호출(엔진 비교 + 정답세트.xlsx 작성)
필요: skyfield, skyfield-data, openpyxl, tzdata(또는 시스템 zoneinfo), astronomy-engine(비교용)
"""
import datetime as dt
import json
import math
import os
import sys

sys.dont_write_bytecode = True  # 상위 폴더(엔진 옆)에 __pycache__를 만들지 않음
from zoneinfo import ZoneInfo

import skyfield
import skyfield_data
from skyfield.api import Loader
from skyfield.framelib import ecliptic_frame

HERE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = os.path.join(HERE, '정답세트.json')
XLSX_PATH = os.path.join(HERE, '정답세트.xlsx')

SIGNS = ['양자리', '황소자리', '쌍둥이자리', '게자리', '사자자리', '처녀자리',
         '천칭자리', '전갈자리', '사수자리', '염소자리', '물병자리', '물고기자리']
BODY_KO = {'sun': '태양', 'moon': '달', 'mercury': '수성', 'venus': '금성', 'mars': '화성',
           'jupiter': '목성', 'saturn': '토성', 'uranus': '천왕성', 'neptune': '해왕성', 'pluto': '명왕성',
           'asc': 'ASC', 'mc': 'MC'}
# DE421 타깃 (목성 이후·명왕성은 바리센터)
DE421_TARGET = {'sun': 'sun', 'moon': 'moon', 'mercury': 'mercury', 'venus': 'venus', 'mars': 'mars',
                'jupiter': 'jupiter barycenter', 'saturn': 'saturn barycenter',
                'uranus': 'uranus barycenter', 'neptune': 'neptune barycenter', 'pluto': 'pluto barycenter'}
PLANETS = list(DE421_TARGET)

PLACES = {
    '서울': (37.5663, 126.9779, 'Asia/Seoul'),
    '부산': (35.1798, 129.0750, 'Asia/Seoul'),
    '제주': (33.4890, 126.4983, 'Asia/Seoul'),
    '경기(수원)': (37.2893, 127.0535, 'Asia/Seoul'),
    '광주': (35.1601, 126.8514, 'Asia/Seoul'),
    '뉴욕': (40.7128, -74.0060, 'America/New_York'),
    '런던': (51.5074, -0.1278, 'Europe/London'),
    '시드니': (-33.8688, 151.2093, 'Australia/Sydney'),
    '로스앤젤레스': (34.0522, -118.2437, 'America/Los_Angeles'),
    '베이징': (39.9042, 116.4074, 'Asia/Shanghai'),
    '도쿄': (35.6762, 139.6503, 'Asia/Tokyo'),
    '호찌민': (10.8231, 106.6297, 'Asia/Ho_Chi_Minh'),
    '모스크바': (55.7558, 37.6173, 'Europe/Moscow'),
}

# (id, 현지시각, 장소, 분류, 메모, fold)
BIRTHS = [
    ('T01', '1965-03-12 07:40', '서울', '일반', '', 0),
    ('T02', '1972-08-25 14:20', '부산', '일반', '', 0),
    ('T03', '1979-12-03 21:10', '제주', '일반', '', 0),
    ('T04', '1983-05-17 03:35', '경기(수원)', '일반', '', 0),
    ('T05', '1991-10-30 11:55', '광주', '일반', '', 0),
    ('T06', '1996-04-01 10:17', '경기(수원)', '일반', '08_샘플_엔진출력과 같은 출생정보', 0),
    ('T07', '2001-02-14 18:45', '부산', '일반', '', 0),
    ('T08', '2005-09-09 06:05', '서울', '일반', '', 0),
    ('T09', '2010-06-30 16:30', '제주', '일반', '', 0),
    ('T10', '1968-11-21 23:20', '광주', '일반', '', 0),
    ('T11', '1987-07-15 14:00', '서울', '한국 서머타임', '1987 서머타임(UTC+10)', 0),
    ('T12', '1988-08-20 09:30', '부산', '한국 서머타임', '1988 서머타임(UTC+10)', 0),
    ('T13', '1987-05-10 01:30', '경기(수원)', '한국 서머타임 경계', '서머타임 시작일 02:00 이전 → 실제 UTC+9. 프로토타입은 날짜 단위라 +10', 0),
    ('T14', '1988-10-09 01:30', '서울', '한국 서머타임 경계', '서머타임 종료일 03:00 이전 → 실제 UTC+10. 프로토타입은 +9', 0),
    ('T15', '1958-07-15 12:00', '서울', '한국 과거 시간대', '표준시 UTC+8:30 + 1958 서머타임 → UTC+9:30. 프로토타입은 +8:30만 처리', 0),
    ('T16', '2000-01-01 00:05', '서울', '자정 경계', '현지 00:05 → UTC 전날', 0),
    ('T17', '1994-03-01 23:55', '광주', '자정 경계', '현지 23:55 (UTC 같은 날 14:55)', 0),
    ('T18', '1993-09-05 18:50', '서울', '달 별자리 경계', '달이 양자리 끝(황소자리 경계 약 0.16° 전)', 0),
    ('T19', '1998-08-10 15:00', '부산', '수성 역행', '수성 역행 한가운데(속도 약 -0.75°/일)', 0),
    ('T20', '2003-03-21 09:30', '서울', '태양 별자리 경계', '춘분(09:59 KST) 약 30분 전 → 태양 물고기자리 29.98°', 0),
    ('T21', '1995-07-04 15:20', '뉴욕', '해외 서머타임', 'EDT UTC-4', 0),
    ('T22', '2001-06-21 08:10', '런던', '해외 서머타임', 'BST UTC+1', 0),
    ('T23', '1999-01-15 19:45', '시드니', '해외 서머타임', 'AEDT UTC+11 (남반구)', 0),
    ('T24', '1990-10-28 01:30', '로스앤젤레스', '해외 모호한 시각', '서머타임 종료로 01:00~02:00이 두 번 → fold=0(첫 번째, PDT UTC-7) 사용. '
            '요청의 1990-11-04는 실제로 모호하지 않아(1990년 종료일은 10-28) 날짜를 바꿈', 0),
    ('T25', '1990-07-01 10:00', '베이징', '해외 서머타임', '중국 서머타임(1986–1991) UTC+9', 0),
    ('T26', '1984-02-29 05:50', '도쿄', '해외', '윤일', 0),
    ('T27', '2007-09-18 13:15', '호찌민', '해외', 'UTC+7', 0),
    ('T28', '1985-06-10 22:40', '모스크바', '고위도', '북위 55.8°, 모스크바 서머타임 UTC+4', 0),
    ('T29', '1960-05-20 08:00', '서울', '한국 과거 시간대', '표준시 UTC+8:30 + 1960 서머타임 → UTC+9:30', 0),
    ('T30', '1977-01-20 04:10', '광주', '일반', '', 0),
]

# 한국 시간대 규칙 추가 점검용 (차트 계산 없이 오프셋만 비교)
TZ_PROBES = [
    '1949-06-01 12:00', '1951-06-01 12:00', '1954-03-20 12:00', '1954-03-21 12:00',
    '1956-06-01 12:00', '1957-06-01 12:00', '1959-06-01 12:00', '1961-08-09 23:50',
    '1961-08-10 01:00', '1987-05-10 01:30', '1987-05-10 03:30', '1987-10-11 01:30',
    '1987-10-11 03:30', '1988-05-08 01:30', '1988-05-08 03:30', '1988-10-09 01:30',
    '1988-10-09 03:30', '1990-06-01 12:00',
]


def to_utc(local_str, tz, fold=0):
    naive = dt.datetime.strptime(local_str, '%Y-%m-%d %H:%M')
    z = ZoneInfo(tz)
    aware = naive.replace(tzinfo=z, fold=fold)
    off0 = naive.replace(tzinfo=z, fold=0).utcoffset()
    off1 = naive.replace(tzinfo=z, fold=1).utcoffset()
    utc = aware.astimezone(dt.timezone.utc)
    back = utc.astimezone(z).replace(tzinfo=None)
    flag = ''
    if off0 != off1:
        flag = '모호한 시각(두 번 존재)' if back == naive else '존재하지 않는 시각(건너뜀)'
    off_h = aware.utcoffset().total_seconds() / 3600
    return utc.replace(tzinfo=None), off_h, flag


def fmt_off(h):
    s = '+' if h >= 0 else '-'
    h = abs(h)
    return f'UTC{s}{int(h)}:{int(round((h % 1) * 60)):02d}'


def wrap180(x):
    return (x + 180) % 360 - 180


class Sky:
    def __init__(self):
        L = Loader(skyfield_data.get_skyfield_data_path(), verbose=False)
        self.ts = L.timescale(builtin=False)  # skyfield-data의 finals2000A.all (ΔT·UT1)
        self.eph = L('de421.bsp')
        self.earth = self.eph['earth']

    def t(self, utc):
        args = (utc.year, utc.month, utc.day, utc.hour, utc.minute, utc.second + utc.microsecond / 1e6)
        # 1972년 이전: skyfield는 UTC를 'TAI−10초'로 가정해 UT1과 최대 ~10초 어긋난다.
        # 당시 상용시(UTC)는 실제로 UT1과 0.1초 이내로 맞춰졌으므로 UT1로 해석한다.
        return self.ts.ut1(*args) if utc.year < 1972 else self.ts.utc(*args)

    def lon(self, body, t):
        app = self.earth.at(t).observe(self.eph[DE421_TARGET[body]]).apparent()
        return float(app.frame_latlon(ecliptic_frame)[1].degrees) % 360

    def chart(self, utc, lat, lng):
        t = self.t(utc)
        out = {}
        h = 1 / 24  # 속도: ±1시간 중앙차분 (사실상 순간 속도)
        for b in PLANETS:
            L0 = self.lon(b, t)
            Lm, Lp = self.lon(b, self.ts.tt_jd(t.tt - h)), self.lon(b, self.ts.tt_jd(t.tt + h))
            speed = wrap180(Lp - Lm) / (2 * h)
            out[b] = {'lon': L0, 'speed_deg_per_day': speed, 'retro': bool(speed < 0)}
        # ASC / MC
        gast = float(t.gast)  # 시간
        lst = (gast * 15 + lng) % 360
        eps = math.degrees(float(t._mean_obliquity_radians + t._nutation_angles_radians[1]))  # 참 황도경사
        r, e, f = math.radians(lst), math.radians(eps), math.radians(lat)
        mc = math.degrees(math.atan2(math.sin(r), math.cos(r) * math.cos(e))) % 360
        asc = math.degrees(math.atan2(math.cos(r), -(math.sin(e) * math.tan(f) + math.cos(e) * math.sin(r)))) % 360
        east = (asc - mc) % 360
        assert 0 < east < 180, f'ASC가 MC 동쪽 180° 안에 있지 않음: {utc} {asc} {mc}'
        # 검산: ASC 지점이 동쪽 지평선(고도 0, 떠오르는 중)에 있는지
        la = math.radians(asc)
        ra = math.atan2(math.sin(la) * math.cos(e), math.cos(la))
        de = math.asin(math.sin(la) * math.sin(e))
        H = r - ra
        alt = math.degrees(math.asin(math.sin(f) * math.sin(de) + math.cos(f) * math.cos(de) * math.cos(H)))
        assert abs(alt) < 1e-6 and math.sin(H) < 0, f'ASC 검산 실패 {utc} alt={alt}'
        out['asc'] = {'lon': asc}
        out['mc'] = {'lon': mc}
        asc_sign = int(asc // 30)
        for k, v in out.items():
            s = int(v['lon'] // 30)
            v['sign'] = SIGNS[s]
            v['deg_in_sign'] = v['lon'] % 30
            v['house'] = (s - asc_sign) % 12 + 1
            v['lon'] = round(v['lon'], 6)
            v['deg_in_sign'] = round(v['deg_in_sign'], 6)
            if 'speed_deg_per_day' in v:
                v['speed_deg_per_day'] = round(v['speed_deg_per_day'], 6)
        extra = {'gast_hours': round(gast, 8), 'lst_deg': round(lst, 6), 'true_obliquity_deg': round(eps, 8),
                 'delta_t_sec': round(float(t.delta_t), 3), 'asc_east_of_mc_deg': round(east, 4)}
        return out, extra


def build():
    sky = Sky()
    births = []
    for bid, local, place, cat, note, fold in BIRTHS:
        lat, lng, tz = PLACES[place]
        utc, off, flag = to_utc(local, tz, fold)
        bodies, extra = sky.chart(utc, lat, lng)
        births.append({
            'id': bid, 'local': local, 'fold': fold, 'place': place, 'lat': lat, 'lng': lng, 'tz': tz,
            'utc_offset_hours': off, 'utc_offset': fmt_off(off), 'utc': utc.strftime('%Y-%m-%dT%H:%M:%SZ'),
            'category': cat, 'note': note, 'tz_flag': flag,
            'reference': {'bodies': bodies, **extra},
        })
    probes = []
    for local in TZ_PROBES:
        utc, off, flag = to_utc(local, 'Asia/Seoul', 0)
        probes.append({'local': local, 'tz': 'Asia/Seoul', 'zoneinfo_offset_hours': off,
                       'utc': utc.strftime('%Y-%m-%dT%H:%M:%SZ'), 'tz_flag': flag})
    meta = {
        'title': '쏠 아스트로 하우스 출생차트 정답 비교 세트',
        'generated': dt.date.today().isoformat(),
        'reference_software': f'skyfield {skyfield.__version__}',
        'ephemeris': 'JPL DE421 (skyfield-data 번들, 1900–2050)',
        'time_scale': 'finals2000A.all (skyfield-data 번들) — UT1/ΔT',
        'frame': '지구 중심 겉보기 위치(광행차·광시간·중력굴절 포함), 그날의 참 황도·참 춘분점(skyfield ecliptic_frame), 회귀황도',
        'targets': '태양·달·수성·금성·화성 = 천체 자체, 목성~명왕성 = 바리센터',
        'retro_rule': '경도 속도(±1시간 중앙차분) < 0 이면 역행',
        'asc_mc': 'GAST(겉보기 항성시)+경도=LST, 참 황도경사. MC=atan2(sinLST, cosLST·cosε), '
                  'ASC=atan2(cosLST, −(sinε·tanφ+cosε·sinLST)). ASC가 MC 동쪽 180° 안·지평선 고도 0인지 검산함',
        'houses': '홀사인: ASC 별자리 = 1하우스',
        'utc_handling': '1972년 이후 UTC(윤초 반영), 1972년 이전은 UTC≈UT1로 해석(skyfield 기본 가정은 최대 10초 어긋남)',
        'timezone': 'IANA zoneinfo(Asia/Seoul 등)를 정답으로 사용. 모호한 시각은 fold=0',
        'latitude': '지리(측지) 위도 사용',
        'tolerances': {'planets_arcsec': 60, 'moon_arcsec': 120, 'asc_mc_deg': 0.1},
    }
    return {'metadata': meta, 'births': births, 'tz_probes': probes}


# ---------------- Excel ----------------
README_LINES = [
    '쏠 아스트로 하우스 — 출생차트 정답 비교 세트',
    '',
    '무엇: 가상의 테스트 출생 30건(T01~T30)에 대해 skyfield + JPL DE421로 계산한 "정답" 값과, 엔진(07_엔진_프로토타입.py, astronomy-engine) 결과 비교.',
    '시트: 출생정보(입력) / 정답(출생×천체 1행) / 엔진비교(차이·PASS/FAIL) / 시간대비교(한국 오프셋 규칙 vs zoneinfo) / 요약(최대 차이·FAIL 수) / 읽어주세요',
    '좌표계: 회귀황도, 지구 중심 겉보기 위치, 그날의 참 황도·춘분점. 하우스는 홀사인. 1972년 이전 시각은 UTC≈UT1로 해석.',
    '허용오차: 행성 ≤ 60″(1′), 달 ≤ 120″, ASC·MC ≤ 0.1°(360″). 별자리·하우스·역행 불일치는 FAIL.',
    '엔진 비교는 UTC를 직접 넣어 계산(엔진의 현지시→UTC 변환은 별도 시간대비교 시트에서 점검).',
    '재생성: python make_reference.py   /   비교만: python compare_engine.py (FAIL이 있으면 종료코드 1)',
    'JSON(정답세트.json)과 이 엑셀은 같은 데이터입니다.',
]


def write_xlsx(data, path=XLSX_PATH):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='DDE7F0')
    fail_fill = PatternFill('solid', fgColor='F8D0D0')

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = 'A2'
        for i, h in enumerate(header, 1):
            w = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(w, 8), 60)
        return ws

    B = data['births']
    sheet('출생정보', ['ID', '현지시각', 'fold', '장소', '위도', '경도', 'IANA 시간대', 'UTC 오프셋', 'UTC',
                     '분류', '메모', '시간대 플래그', 'GAST(시)', 'LST(°)', '참 황도경사(°)', 'ΔT(초)'],
          [[b['id'], b['local'], b['fold'], b['place'], b['lat'], b['lng'], b['tz'], b['utc_offset'], b['utc'],
            b['category'], b['note'], b['tz_flag'], b['reference']['gast_hours'], b['reference']['lst_deg'],
            b['reference']['true_obliquity_deg'], b['reference']['delta_t_sec']] for b in B], first=True)
    rows = []
    for b in B:
        for k, v in b['reference']['bodies'].items():
            rows.append([b['id'], k, BODY_KO[k], v['lon'], v['sign'], v['deg_in_sign'], v['house'],
                         v.get('speed_deg_per_day', ''), ('역행' if v['retro'] else '순행') if 'retro' in v else ''])
    sheet('정답', ['ID', '천체', '천체(한글)', '황경(°)', '별자리', '별자리 내 도수(°)', '홀사인 하우스',
                 '속도(°/일)', '역행'], rows)
    comp = data.get('engine_comparison')
    if comp:
        ws = sheet('엔진비교', ['ID', '천체', '정답 황경(°)', '엔진 황경(°)', '차이(″)', '허용(″)', '정답 별자리',
                             '엔진 별자리', '별자리 일치', '정답 하우스', '엔진 하우스', '하우스 일치', '정답 역행',
                             '엔진 역행', '역행 일치', '판정', '비고'],
                   [[r['id'], r['body'], r['ref_lon'], r['engine_lon'], r['diff_arcsec'], r['tol_arcsec'],
                     r['ref_sign'], r['engine_sign'], r['sign_match'], r['ref_house'], r['engine_house'],
                     r['house_match'], r['ref_retro'], r['engine_retro'], r['retro_match'], r['result'], r['note']]
                    for r in comp['rows']])
        for row in ws.iter_rows(min_row=2):
            if row[15].value == 'FAIL':
                for c in row:
                    c.fill = fail_fill
        ws2 = sheet('시간대비교', ['구분', 'ID', '현지시각', 'zoneinfo 오프셋(시)', '프로토타입 규칙 오프셋(시)',
                                '차이(분)', '일치', 'zoneinfo 플래그'],
                    [[r['kind'], r['id'], r['local'], r['zoneinfo_offset_hours'], r['prototype_offset_hours'],
                      r['diff_minutes'], r['match'], r['tz_flag']] for r in comp['tz_rows']])
        for row in ws2.iter_rows(min_row=2):
            if row[6].value is False:
                for c in row:
                    c.fill = fail_fill
        s = comp['summary']
        sum_rows = [[k, s['max_abs_diff_arcsec'][k], s['tol_arcsec'][k]] for k in s['max_abs_diff_arcsec']]
        ws3 = sheet('요약', ['천체', '최대 |차이|(″)', '허용(″)'], sum_rows)
        ws3.append([])
        for k in ('fail_count', 'sign_mismatch', 'house_mismatch', 'retro_mismatch', 'tz_mismatch_births', 'tz_mismatch_probes'):
            v = s[k]
            ws3.append([k, len(v) if isinstance(v, list) else v] + (v if isinstance(v, list) else []))
        ws3.append(['warn_count', s['warn_count']])
        ws3.append(['engine', s['engine']])
        ws3.append(['run_date', s['run_date']])
    ws = wb.create_sheet('읽어주세요')
    ws.column_dimensions['A'].width = 130
    for line in README_LINES:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=13)
    meta = data['metadata']
    ws.append([])
    for k, v in meta.items():
        ws.append([f'{k}: {json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v}'])
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True)
    wb.move_sheet('읽어주세요', offset=-(len(wb.sheetnames) - 1))
    wb.save(path)


def save_json(data, path=JSON_PATH):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    data = build()
    save_json(data)
    print(f'정답 {len(data["births"])}건 저장: {JSON_PATH}')
    sys.path.insert(0, HERE)
    import compare_engine
    sys.exit(compare_engine.run(data))
